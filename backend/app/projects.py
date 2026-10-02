"""Uploaded projects: parsing, validation, and the project shape the rest of the backend consumes.

Nothing here alters analyst data silently: rows with missing values are dropped only when the analyst
explicitly asks for listwise deletion, and the count is recorded."""

from __future__ import annotations

import csv
import hashlib
import io
import re
import uuid
from typing import Any, cast

from app.db.models import ProjectRow
from app.demo import DEMO_DIR, load_demo
from app.domain.prompt import VariableSpec
from app.domain.stats import describe

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 50_000
MAX_CONDITIONS = 8
NAME_RE = re.compile(r"^[A-Za-z][A-Za-z0-9_]*$")
ANCHOR_KEYS = ("full_non_membership", "crossover", "full_membership")
BREAK_KEYS = ("break_0", "break_33", "break_67")
KINDS = ("direct", "breakpoints", "precalibrated")


class ProjectError(ValueError):
    """User-facing problem with an upload or configuration."""


def _num(x: str | None) -> float | None:
    if x is None or x.strip() == "":
        return None
    v = float(x.strip().replace(",", "")) if re.fullmatch(r"\s*-?\d{1,3}(,\d{3})+(\.\d+)?\s*", x) else float(x.strip())
    if v != v or v in (float("inf"), float("-inf")):
        raise ValueError("non-finite")
    return v


def parse_upload(filename: str, content: bytes) -> tuple[list[str], list[list[str | None]]]:
    """Returns (header, rows of strings). CSV (comma, semicolon or tab) or XLSX (first sheet)."""
    if len(content) > MAX_BYTES:
        raise ProjectError("file is larger than 10 MB")
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            ws = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
        except Exception as e:
            raise ProjectError(f"could not read the Excel file: {e}") from e
        raw = [["" if c is None else str(c) for c in row] for row in ws.iter_rows(values_only=True)]
    elif name.endswith((".csv", ".tsv", ".txt")):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("latin-1")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        raw = [row for row in csv.reader(io.StringIO(text), dialect)]
    else:
        raise ProjectError("unsupported file type; upload .csv or .xlsx")
    raw = [r for r in raw if any(c.strip() for c in r)]
    if len(raw) < 2:
        raise ProjectError("the file needs a header row and at least one data row")
    header = [h.strip() for h in raw[0]]
    if any(not h for h in header):
        raise ProjectError("every column needs a header name")
    if len(set(header)) != len(header):
        raise ProjectError("column names must be unique")
    body = raw[1:]
    if len(body) > MAX_ROWS:
        raise ProjectError(f"more than {MAX_ROWS} rows")
    width = len(header)
    rows: list[list[str | None]] = [[(c.strip() if c.strip() != "" else None) for c in (r + [""] * width)[:width]] for r in body]
    return header, rows


def column_info(header: list[str], rows: list[list[str | None]]) -> list[dict[str, Any]]:
    out = []
    for j, h in enumerate(header):
        vals = [r[j] for r in rows]
        present = [v for v in vals if v is not None]
        numeric = bool(present)
        for v in present:
            try:
                _num(v)
            except ValueError:
                numeric = False
                break
        nums = [cast(float, _num(v)) for v in present] if numeric else []
        out.append({"name": h, "numeric": numeric, "n_missing": len(vals) - len(present), "n_unique": len(set(present)),
                    "min": min(nums) if nums else None, "max": max(nums) if nums else None})
    return out


def _to_csv(header: list[str], rows: list[list[str | None]]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(header)
    w.writerows([["" if c is None else c for c in r] for r in rows])
    return buf.getvalue()


def _from_csv(text: str) -> tuple[list[str], list[list[str | None]]]:
    r = list(csv.reader(io.StringIO(text)))
    return r[0], [[c if c != "" else None for c in row] for row in r[1:]]


def create_upload(session: Any, filename: str, content: bytes) -> dict[str, Any]:
    header, rows = parse_upload(filename, content)
    pid = "p_" + uuid.uuid4().hex[:10]
    text = _to_csv(header, rows)
    session.add(ProjectRow(id=pid, filename=filename, name=filename.rsplit(".", 1)[0],
                           dataset_sha256=hashlib.sha256(text.encode()).hexdigest(), dataset_csv=text))
    session.commit()
    return {"project_id": pid, "filename": filename, "n_rows": len(rows), "columns": column_info(header, rows),
            "preview": {"header": header, "rows": rows[:8]}}


def validate_config(cfg: dict[str, Any], info: dict[str, dict[str, Any]], n_rows: int) -> list[str]:
    """Returns a list of problems (empty = ok)."""
    errs: list[str] = []
    vs = cfg.get("variables", [])
    conds = [v for v in vs if v["role"] == "condition"]
    outs = [v for v in vs if v["role"] == "outcome"]
    if len(outs) != 1:
        errs.append("choose exactly one outcome")
    if len(conds) < 2:
        errs.append("choose at least two conditions")
    if len(conds) > MAX_CONDITIONS:
        errs.append(f"at most {MAX_CONDITIONS} conditions are supported")
    for v in vs:
        n = v["name"]
        if n not in info:
            errs.append(f"{n}: no such column")
            continue
        if not info[n]["numeric"]:
            errs.append(f"{n}: column is not numeric")
        if not NAME_RE.match(n):
            errs.append(f"{n}: QCA variable names must start with a letter and use only letters, digits and underscores; rename the column in your file")
        if v["direction"] not in ("positive", "negative"):
            errs.append(f"{n}: direction must be positive or negative")
        kind = v.get("calibration", "direct")
        if kind not in KINDS:
            errs.append(f"{n}: calibration must be one of {', '.join(KINDS)}")
            continue
        if kind == "precalibrated":
            if v["role"] == "outcome":
                errs.append(f"{n}: the outcome must be calibrated by the agents (anchors or breakpoints); only conditions can be taken as already calibrated")
            if info[n].get("min") is not None and (info[n]["min"] < 0 or info[n]["max"] > 1):
                errs.append(f"{n}: an already-calibrated variable must hold membership scores between 0 and 1 (found {info[n]['min']:g} to {info[n]['max']:g})")
            if v.get("anchors") or v.get("breakpoints"):
                errs.append(f"{n}: an already-calibrated variable takes no anchors or breakpoints")
            if info[n]["n_unique"] < 2:
                errs.append(f"{n}: a constant variable cannot be used")
        elif info[n]["n_unique"] < 3:
            errs.append(f"{n}: fewer than 3 distinct values cannot be calibrated")
        if kind == "direct":
            a = v.get("anchors")
            if a:
                if any(a.get(k) is None for k in ANCHOR_KEYS):
                    errs.append(f"{n}: original anchors must have all three values or none")
                else:
                    fn, cr, fm = (float(a[k]) for k in ANCHOR_KEYS)
                    ok = fn < cr < fm if v["direction"] == "positive" else fn > cr > fm
                    if not ok:
                        errs.append(f"{n}: original anchors are not ordered for a {v['direction']} variable")
        if kind == "breakpoints":
            b = v.get("breakpoints")
            if v.get("anchors"):
                errs.append(f"{n}: a breakpoint variable takes breakpoints, not anchors")
            if b:
                if any(b.get(k) is None for k in BREAK_KEYS):
                    errs.append(f"{n}: original breakpoints must have all three values or none")
                else:
                    b0, b1, b2 = (float(b[k]) for k in BREAK_KEYS)
                    ok = b0 < b1 < b2 if v["direction"] == "positive" else b0 > b1 > b2
                    if not ok:
                        errs.append(f"{n}: original breakpoints are not ordered for a {v['direction']} variable")
        if v["role"] == "condition" and v.get("dir_exp") not in (None, 0, 1):
            errs.append(f"{n}: directional expectation must be 1, 0 or empty")
    if len({v["name"] for v in vs}) != len(vs):
        errs.append("a column was selected twice")
    miss = sum(info[v["name"]]["n_missing"] for v in vs if v["name"] in info)
    if miss and not cfg.get("drop_missing"):
        errs.append("selected columns contain missing values; tick 'drop rows with missing values' to continue")
    tt = cfg.get("reference_cutoffs")
    if tt:
        c = tt.get("consistency_threshold")
        if c is not None and not 0 <= c <= 1:
            errs.append("consistency cutoff must be between 0 and 1")
        f = tt.get("frequency_threshold")
        if f is not None and (f < 1 or int(f) != f):
            errs.append("frequency cutoff must be a whole number of at least 1")
    return errs


def configure(session: Any, pid: str, cfg: dict[str, Any]) -> dict[str, Any]:
    row = session.get(ProjectRow, pid)
    if row is None:
        raise ProjectError("unknown project")
    header, rows = _from_csv(row.dataset_csv)
    info = {c["name"]: c for c in column_info(header, rows)}
    errs = validate_config(cfg, info, len(rows))
    if errs:
        raise ProjectError("; ".join(errs))
    row.name = cfg.get("name") or row.name
    row.config = cfg
    session.commit()
    return project_summary(get_project(session, pid))


def _build(name: str, case_description: str, header: list[str], rows: list[list[str | None]], cfg: dict[str, Any]) -> dict[str, Any]:
    vs = cfg["variables"]
    idx = {h: i for i, h in enumerate(header)}
    keep = [r for r in rows if all(r[idx[v["name"]]] is not None for v in vs)]
    data = {v["name"]: [cast(float, _num(r[idx[v["name"]]])) for r in keep] for v in vs}
    variables = [VariableSpec(v["name"], v["role"], "positive" if v.get("calibration") == "precalibrated" else v["direction"],
                              v.get("construct_definition", ""), v.get("instrument", ""), v.get("units", ""),
                              describe(data[v["name"]]), v.get("calibration", "direct")) for v in vs]
    kinds = {v["name"]: v.get("calibration", "direct") for v in vs}
    reference: dict[str, dict[str, float]] = {}
    for v in vs:
        k = kinds[v["name"]]
        if k == "direct" and v.get("anchors"):
            reference[v["name"]] = {key: float(v["anchors"][key]) for key in ANCHOR_KEYS}
        elif k == "breakpoints" and v.get("breakpoints"):
            reference[v["name"]] = {key: float(v["breakpoints"][key]) for key in BREAK_KEYS}
    needed = [v["name"] for v in vs if kinds[v["name"]] != "precalibrated"]
    tt = cfg.get("reference_cutoffs") or {}
    has_ref = (all(n in reference for n in needed) and tt.get("consistency_threshold") is not None
               and tt.get("frequency_threshold") is not None)
    outcome = next(v["name"] for v in vs if v["role"] == "outcome")
    return {
        "id": None, "name": name, "project": {"name": name, "case_description": case_description}, "data": data,
        "variables": variables, "outcome": outcome,
        "directions": {v["name"]: ("positive" if v.get("calibration") == "precalibrated" else v["direction"]) for v in vs},
        "kinds": kinds, "condition_order": [v["name"] for v in vs if v["role"] == "condition"],
        "dir_exp": {v["name"]: v.get("dir_exp") for v in vs if v["role"] == "condition"},
        "reference": reference if has_ref else {}, "reference_cutoffs": tt if has_ref else {},
        "has_reference": has_ref, "n_cases": len(keep), "n_dropped": len(rows) - len(keep),
    }


def get_project(session: Any, pid: str) -> dict[str, Any]:
    if pid == "demo":
        d = load_demo()
        d.update({"dataset_sha256": hashlib.sha256(((DEMO_DIR / "dataset.csv")).read_bytes()).hexdigest(), "id": "demo", "name": d["project"]["name"], "has_reference": True, "n_cases": len(next(iter(d["data"].values()))), "n_dropped": 0})
        return d
    row = session.get(ProjectRow, pid)
    if row is None:
        raise ProjectError("unknown project")
    if not row.config:
        raise ProjectError("project is not configured yet")
    header, rows = _from_csv(row.dataset_csv)
    d = _build(row.name, row.config.get("case_description", ""), header, rows, row.config)
    d["id"] = pid
    d["dataset_sha256"] = row.dataset_sha256
    return d


def project_summary(d: dict[str, Any]) -> dict[str, Any]:
    from app.domain.prompt import case_description_warnings, prompt_warnings

    return {
        "id": d["id"], "name": d["name"], "case_description": d["project"]["case_description"],
        "variables": [v.__dict__ for v in d["variables"]], "reference": d["reference"],
        "reference_cutoffs": d["reference_cutoffs"], "dir_exp": d["dir_exp"], "has_reference": d["has_reference"],
        "warnings": [*prompt_warnings(d["variables"]), *case_description_warnings(d["project"]["case_description"])],
        "n_cases": d["n_cases"], "n_dropped": d["n_dropped"],
        "is_demo": d["id"] == "demo",
    }


def get_setup(session: Any, pid: str) -> dict[str, Any]:
    """Everything the setup form needs to reopen an uploaded project for editing."""
    if pid == "demo":
        raise ProjectError("the demo project cannot be edited; upload your own data")
    row = session.get(ProjectRow, pid)
    if row is None:
        raise ProjectError("unknown project")
    header, rows = _from_csv(row.dataset_csv)
    return {"upload": {"project_id": pid, "filename": row.filename, "n_rows": len(rows), "columns": column_info(header, rows),
                       "preview": {"header": header, "rows": rows[:8]}}, "config": row.config}
